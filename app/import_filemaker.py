#!/usr/bin/env python3
"""
One-time FileMaker data import.

Usage (inside the container):
  docker compose exec app python import_filemaker.py

Or with a custom path:
  docker compose exec app python import_filemaker.py /filemaker-import/export\ filemaker
"""

import asyncio
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from sqlalchemy import select

from database import AsyncSessionLocal, engine, Base
from models import Group, Subgroup, PendingScan, Document

IMPORT_DIR = Path(sys.argv[1] if len(sys.argv) > 1 else "/filemaker-import/export filemaker")


def _num(val) -> int | None:
    """Extract the numeric portion from values like 'M83', 83, '83'."""
    if val is None:
        return None
    m = re.search(r'\d+', str(val))
    return int(m.group()) if m else None


def read_xlsx(filename: str) -> tuple[dict, list]:
    """Return (header_index, data_rows). Skips blank rows."""
    wb = openpyxl.load_workbook(IMPORT_DIR / filename)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_index = {h: i for i, h in enumerate(rows[0])}
    data = [list(r) for r in rows[1:] if any(v is not None for v in r)]
    return header_index, data


async def run():
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    g_idx, g_rows = read_xlsx("Groups.tab.xlsx")
    s_idx, s_rows = read_xlsx("Subgroups.xlsx")
    p_idx, p_rows = read_xlsx("Phone.xlsx")
    f_idx, f_rows = read_xlsx("Files.xlsx")

    async with AsyncSessionLocal() as db:

        # ------------------------------------------------------------------
        # Groups
        # ------------------------------------------------------------------
        print("Importing groups…")
        fm_group_num_to_id: dict[int, int] = {}   # #Groupnumonly → Group.id
        fm_group_text_to_id: dict[str, int] = {}  # #GroupID text → Group.id

        for row in g_rows:
            name = row[g_idx['Group name']]
            if not name:
                continue
            include = bool(row[g_idx['Include group name in path by default']])
            fm_text = str(row[g_idx['#GroupID']]).strip()
            fm_num = _num(row[g_idx['#Groupnumonly']])

            res = await db.execute(select(Group).where(Group.name == name))
            g = res.scalar_one_or_none()
            if not g:
                g = Group(name=name, include_in_path=include)
                db.add(g)
                await db.flush()
                print(f"  + {name}")
            else:
                print(f"  = {name} (exists)")

            if fm_num:
                fm_group_num_to_id[fm_num] = g.id
            fm_group_text_to_id[fm_text] = g.id

        await db.commit()

        # ------------------------------------------------------------------
        # Subgroups  (export has one row per subgroup + extra rows for
        # related files; skip rows where #SubgroupID is None)
        # ------------------------------------------------------------------
        print("\nImporting subgroups…")
        fm_sub_num_to_id: dict[int, int] = {}  # numeric subgroup ID → Subgroup.id

        for row in s_rows:
            fm_sub_text = row[s_idx['#SubgroupID']]
            if fm_sub_text is None:
                continue  # related-record rows

            sub_name = row[s_idx['Subgroup name']]
            if not sub_name:
                continue

            fm_group_text = str(row[s_idx['GroupID']]).strip() if row[s_idx['GroupID']] else None
            group_id = fm_group_text_to_id.get(fm_group_text)
            if not group_id:
                print(f"  ! Unknown group '{fm_group_text}' for subgroup '{sub_name}' — skip")
                continue

            filename_tag = row[s_idx['Filename Tag']] or None
            custom_path = row[s_idx['Custom Path']] or None
            fm_num = _num(fm_sub_text)

            res = await db.execute(
                select(Subgroup).where(Subgroup.group_id == group_id, Subgroup.name == sub_name)
            )
            s = res.scalar_one_or_none()
            if not s:
                s = Subgroup(
                    group_id=group_id,
                    name=sub_name,
                    # Don't store tag if it's the same as the name or a '/' placeholder
                    filename_tag=filename_tag if (filename_tag and filename_tag != sub_name and filename_tag != '/') else None,
                    custom_path=custom_path,
                )
                db.add(s)
                await db.flush()
                print(f"  + {sub_name} (group_id={group_id})")
            else:
                print(f"  = {sub_name} (exists)")

            if fm_num is not None:
                fm_sub_num_to_id[fm_num] = s.id

        await db.commit()

        # ------------------------------------------------------------------
        # Pending scans (On Phone)
        # All historical records are already filed, so status = 'filed'
        # ------------------------------------------------------------------
        print("\nImporting phone records…")
        fm_phone_num_to_id: dict[int, int] = {}

        for row in p_rows:
            fm_id_text = row[p_idx['#']]
            if fm_id_text is None:
                continue
            fm_num = row[p_idx['#numonly']]
            if fm_num is None:
                continue
            fm_num = int(fm_num)

            sub_fm_num = row[p_idx['SubgroupID']]
            subgroup_id = fm_sub_num_to_id.get(int(sub_fm_num)) if sub_fm_num is not None else None

            created_at = row[p_idx['Item Created']]
            if not isinstance(created_at, datetime):
                created_at = datetime.utcnow()

            p = PendingScan(
                subgroup_id=subgroup_id,
                item_one_off_description=row[p_idx['Item one off description']] or None,
                created_at=created_at,
                rob_task=bool(row[p_idx['Rob Task Flag']]),
                beccy_task=bool(row[p_idx['Beccy Task Flag']]),
                task_note=row[p_idx['Task Note']] or None,
                comment=row[p_idx['Comment']] or None,
                status="filed",
            )
            db.add(p)
            await db.flush()
            fm_phone_num_to_id[fm_num] = p.id

        await db.commit()
        print(f"  {len(fm_phone_num_to_id)} phone records imported")

        # ------------------------------------------------------------------
        # Documents (Files)
        # ------------------------------------------------------------------
        print("\nImporting file records…")
        count = 0

        for row in f_rows:
            if row[f_idx['#']] is None:
                continue

            google_file_id = row[f_idx['Google FileID']] or None
            if google_file_id:
                res = await db.execute(select(Document).where(Document.drive_file_id == google_file_id))
                if res.scalar_one_or_none():
                    continue

            fm_group_num = row[f_idx['Assigned | Group ID']]
            fm_sub_num = row[f_idx['Assigned | Subgroup ID']]
            fm_phone_num = row[f_idx['Assigned | Phone ID #']]

            ts = row[f_idx['Timestamp from filename at import']]
            moved_ts = row[f_idx['Moved and renamed timestamp']]
            imported_at = row[f_idx['Time file was imported']]

            doc = Document(
                drive_file_id=google_file_id,
                file_name_when_imported=row[f_idx['File Name when imported']] or '',
                timestamp_from_filename=ts if isinstance(ts, datetime) else None,
                time_file_was_imported=imported_at if isinstance(imported_at, datetime) else datetime.utcnow(),
                group_id=fm_group_num_to_id.get(int(fm_group_num)) if fm_group_num else None,
                subgroup_id=fm_sub_num_to_id.get(int(fm_sub_num)) if fm_sub_num else None,
                pending_scan_id=fm_phone_num_to_id.get(int(fm_phone_num)) if fm_phone_num else None,
                assigned_item_one_off_description=row[f_idx['Assigned | Item one off description']] or None,
                assigned_comment=row[f_idx['Assigned | Comments']] or None,
                assigned_rob_task=bool(row[f_idx['Assigned | Rob Task Flag']]),
                assigned_beccy_task=bool(row[f_idx['Assigned | Beccy Task Flag']]),
                assigned_task_note=row[f_idx['Assigned | Task Note']] or None,
                desired_folder_location=row[f_idx['Desired folder location']] or None,
                desired_filename=row[f_idx['Desired Filename']] or None,
                desired_full_path=row[f_idx['Desired full path']] or None,
                current_full_path=row[f_idx['Current full path']] or None,
                moved_and_renamed_timestamp=moved_ts if isinstance(moved_ts, datetime) else None,
                broken_link=bool(row[f_idx['Broken link']]),
            )
            db.add(doc)
            count += 1
            if count % 100 == 0:
                await db.flush()
                print(f"  {count} files…")

        await db.commit()
        print(f"  {count} file records imported")
        print("\nAll done.")


if __name__ == "__main__":
    asyncio.run(run())
